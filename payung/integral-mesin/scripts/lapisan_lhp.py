"""Lapisan LHP — padanan lapisan aplikasi server INTEGRAL (backend/app/tools/lhr_tools.py).

`render_lhp.py` V6 menghasilkan LHP berparadigma REVIU. Di server, lapisan aplikasi
lalu (A1) menyesuaikan judul/kata dan nama berkas per rumpun jenis (LHA/LHR/LHE/LP),
(A2) mengganti paragraf metodologi/intro/simpulan sesuai jenis, dan memakai perender
tersendiri untuk profil memo (konsultansi), pendampingan (konsultasi PBJ), dan
rb-4dim (evaluasi RB). Di Cowork tak ada server — lapisan itu dibawa ke sini.

Diturunkan dari lhr_tools.py (29 Sep 2026). Satu penyempurnaan yang DISENGAJA:
pernyataan kesesuaian SAIPI 2430 ("dilaksanakan sesuai dengan Standar Audit Intern
Pemerintah Indonesia") dipastikan ada di SEMUA profil. Di server FULL kalimat itu
hanya ada untuk rumpun audit dan berbunyi "sesuai Standar…" (tanpa "dengan"), sehingga
QC KOM-008 menandai KRITIS setiap LHP — terbukti pada LHR server hasil E2E 27 Sep.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document

PERNYATAAN_SAIPI = ("Penugasan ini dilaksanakan sesuai dengan Standar Audit Intern Pemerintah Indonesia "
                    "(SAIPI, PER-01/AAIPI/DPN/2021).")
_POLA_SAIPI = re.compile(r"sesuai\s+dengan\s+Standar\s+Audit\s+Intern\s+Pemerintah\s+Indonesia", re.I)
_BAWAAN_PROFIL = {"konsultansi-umum": "memo", "konsultasi-pengadaan": "pendampingan",
                  "evaluasi-reformasi-birokrasi": "rb-4dim"}
_PROFIL_SAH = {"kksa", "memo", "rb-4dim", "pendampingan"}

_JENIS_LABEL = {
    "audit": ("AUDIT", "Audit", "LHA"), "kepatuhan": ("AUDIT", "Audit", "LHA"),
    "evaluasi": ("EVALUASI", "Evaluasi", "LHE"), "pemantauan": ("PEMANTAUAN", "Pemantauan", "LP"),
    "reviu": ("REVIU", "Reviu", "LHR"),
}
_SUBSTANCE = {
    "audit": {
        "desk review": "Audit dilaksanakan sesuai dengan Standar Audit Intern Pemerintah Indonesia (SAIPI) melalui penelaahan dokumen, pengujian bukti secara memadai, klarifikasi/wawancara kepada pihak terkait, dan analisis untuk memperoleh keyakinan memadai atas hal yang diaudit.",
        "dikelompokkan ke dalam": "Berdasarkan pengujian atas dokumen dan bukti audit, tim Inspektorat II memperoleh {n} ({nt}) temuan yang dikelompokkan ke dalam {m} ({mt}) aspek sesuai sasaran audit.",
        "limited assurance": "Berdasarkan hasil audit yang kami lakukan dengan tingkat keyakinan memadai, terdapat {n} ({nt}) temuan yang perlu ditindaklanjuti auditi sesuai rekomendasi.",
    },
    "evaluasi": {
        "desk review": "Evaluasi dilaksanakan sesuai dengan Standar Audit Intern Pemerintah Indonesia (SAIPI) melalui penelaahan dokumen, analisis data kinerja/capaian, dan klarifikasi kepada unit terkait, dengan membandingkan kondisi yang ada terhadap kriteria evaluasi.",
        "dikelompokkan ke dalam": "Berdasarkan penelaahan, tim Inspektorat II memperoleh {n} ({nt}) catatan evaluasi yang dikelompokkan ke dalam {m} ({mt}) aspek sesuai sasaran.",
        "limited assurance": "Berdasarkan hasil evaluasi dengan tingkat keyakinan terbatas, terdapat {n} ({nt}) catatan yang perlu ditindaklanjuti untuk meningkatkan kualitas penyelenggaraan.",
    },
    "pemantauan": {
        "desk review": "Pemantauan dilaksanakan sesuai dengan Standar Audit Intern Pemerintah Indonesia (SAIPI) melalui penelaahan laporan berkala dan data status pelaksanaan dari auditi/pengawas pekerjaan, dengan membandingkan realisasi terhadap rencana.",
        "dikelompokkan ke dalam": "Berdasarkan pemantauan, tim Inspektorat II mencatat {n} ({nt}) isu/kondisi yang perlu perhatian, dikelompokkan ke dalam {m} ({mt}) aspek.",
        "limited assurance": "Berdasarkan hasil pemantauan, terdapat {n} ({nt}) kondisi yang perlu perhatian sebagaimana diuraikan. Laporan ini bersifat pelaporan status — tidak memberikan keyakinan.",
    },
}


def profil(akar: Path, jenis: str) -> str:
    md = akar / "skills" / jenis / "SKILL.md"
    if md.exists():
        m = re.search(r"^format_laporan:\s*(\S+)", md.read_text(encoding="utf-8"), re.M)
        if m and m.group(1).strip() in _PROFIL_SAH:
            return m.group(1).strip()
    return _BAWAAN_PROFIL.get(jenis, "kksa")


def _meta(jenis: str) -> tuple[str, str, str]:
    for k, v in _JENIS_LABEL.items():
        if jenis.startswith(k):
            return v
    return ("PENGAWASAN", "Pengawasan", "LHP")


def _rumpun(jenis: str) -> str:
    for k in ("audit", "kepatuhan", "evaluasi", "pemantauan", "reviu"):
        if jenis.startswith(k):
            return "audit" if k == "kepatuhan" else k
    return "lain"


def _terbilang(n: int) -> str:
    d = ["nol", "satu", "dua", "tiga", "empat", "lima", "enam", "tujuh", "delapan", "sembilan", "sepuluh", "sebelas"]
    return d[n] if 0 <= n < len(d) else str(n)


def _hitung(d: Path) -> tuple[int, int]:
    n = m = 0
    try:
        n = len(json.loads((d / "_KKP" / "temuan.json").read_text(encoding="utf-8")).get("temuan", []))
    except (OSError, json.JSONDecodeError):
        pass
    try:
        m = len(json.loads((d / "_PKP" / "sasaran-assignment.json").read_text(encoding="utf-8")).get("sasaran", []))
    except (OSError, json.JSONDecodeError):
        pass
    return n, m


def _set(p, teks: str) -> None:
    if p.runs:
        p.runs[0].text = teks
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = teks


def _ganti_kata(p, subs) -> None:
    full = "".join(r.text for r in p.runs)
    if not full:
        return
    new = full
    for pat, repl in subs:
        new = pat.sub(repl, new)
    if new != full:
        _set(p, new)


def _paragraf_semua(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def pastikan_saipi(docx: Path) -> bool:
    """Tambahkan pernyataan SAIPI 2430 bila belum ada. Return True bila ditambahkan."""
    doc = Document(str(docx))
    if any(_POLA_SAIPI.search(p.text) for p in _paragraf_semua(doc)):
        return False
    sasaran = next((p for p in doc.paragraphs if "desk review" in p.text.lower()), None) \
        or next((p for p in doc.paragraphs if re.search(r"metodologi", p.text, re.I)), None)
    if sasaran is not None:
        sasaran.add_run(" " + PERNYATAAN_SAIPI)
    else:
        doc.add_paragraph(PERNYATAAN_SAIPI)
    doc.save(str(docx))
    return True


class BabKurang(ValueError):
    """Kerangka tak memuat bab wajib yang hanya bisa diisi skill (mis. Simpulan)."""


def _ada(doc, pola: str) -> bool:
    r = re.compile(pola, re.I)
    return any(r.search(p.text) for p in _paragraf_semua(doc))


def lengkapi_bab(docx: Path, d: Path, simpulan: str | None) -> list[str]:
    """Tambah bab wajib QC (Tujuan/Sasaran, Ruang Lingkup, Simpulan) bila kerangka tak memuatnya."""
    doc = Document(str(docx))
    ditambah = []
    tj, rl = _tujuan_lingkup(d)
    if not _ada(doc, r"\b(Tujuan|Sasaran)\b"):
        doc.add_heading("Tujuan dan Sasaran", level=1); doc.add_paragraph(tj or "[DIISI AUDITOR — Tujuan dari DPP]"); ditambah.append("Tujuan")
    if not _ada(doc, r"Ruang\s+Lingkup"):
        doc.add_heading("Ruang Lingkup", level=1); doc.add_paragraph(rl or "[DIISI AUDITOR — Ruang Lingkup dari DPP]"); ditambah.append("Ruang Lingkup")
    if not _ada(doc, r"\b(Simpulan|Kesimpulan)\b"):
        if not (simpulan or "").strip():
            raise BabKurang("kerangka LHP ini tidak memuat bab Simpulan — berikan --simpulan (2–4 kalimat, bahasa keyakinan sesuai jenis)")
        doc.add_heading("Simpulan", level=1); doc.add_paragraph(simpulan.strip()); ditambah.append("Simpulan")
    doc.save(str(docx))
    return ditambah


def isi_bagian(docx: Path, isian: dict) -> tuple[list[str], list[str]]:
    """Ganti penanda '[DIISI — …]' dengan isi dari skill. Kunci = awal teks penanda
    (mis. "[DIISI — Komposisi tim"); nilai = teks (paragraf dipisah baris kosong) atau
    {"tabel": [[kepala…], [baris…], …]}. Return (terisi, kunci tak ditemukan)."""
    from copy import deepcopy
    doc = Document(str(docx))
    terisi, dipakai = [], set()
    for p in list(_paragraf_semua(doc)):
        for m in re.findall(r"\[DIISI[^\]]*\]", p.text):
            kunci = next((k for k in isian if m.startswith(k) or k in m), None)
            if kunci is None:
                continue
            nilai = isian[kunci]
            dipakai.add(kunci)
            if isinstance(nilai, dict) and "tabel" in nilai:
                baris = nilai["tabel"]
                _set(p, p.text.replace(m, "").strip())
                t = doc.add_table(rows=len(baris), cols=max(len(b) for b in baris))
                try:
                    t.style = "Table Grid"
                except (KeyError, ValueError):
                    pass
                for i, b in enumerate(baris):
                    for j, v in enumerate(b):
                        t.cell(i, j).text = str(v)
                p._p.addnext(t._tbl)
            else:
                bagian = [x.strip() for x in str(nilai).split("\n\n") if x.strip()]
                _set(p, p.text.replace(m, bagian[0] if bagian else ""))
                acuan = p._p
                for x in bagian[1:]:
                    baru = deepcopy(p._p)
                    for r in baru.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
                        r.text = ""
                    acuan.addnext(baru); acuan = baru
                    from docx.text.paragraph import Paragraph
                    _set(Paragraph(baru, p._parent), x)
            terisi.append(m[:60])
    doc.save(str(docx))
    return terisi, sorted(set(isian) - dipakai)


def finalisasi_kksa(d: Path, jenis: str, simpulan: str | None = None) -> Path | None:
    """A1 + A2 + SAIPI 2430 atas LHP-SUBSTANSI*.docx terbaru; ganti nama per rumpun."""
    keluaran = sorted((d / "_LHP").glob("LHP-SUBSTANSI*.docx"), key=lambda p: p.stat().st_mtime)
    if not keluaran:
        return None
    docx = keluaran[-1]
    up, title, awalan = _meta(jenis)
    rumpun = _rumpun(jenis)
    if rumpun != "reviu":
        n, m = _hitung(d)
        fmt = {"n": n, "nt": _terbilang(n), "m": m, "mt": _terbilang(m)}
        subst = _SUBSTANCE.get(rumpun, {})
        subs = [(re.compile(r"\bREVIU\b"), up), (re.compile(r"\bReviu\b"), title), (re.compile(r"\breviu\b"), title.lower())]
        doc = Document(str(docx))
        for p in _paragraf_semua(doc):
            for penanda, tpl in subst.items():
                if penanda in p.text:
                    _set(p, tpl.format(**fmt))
                    break
            else:
                _ganti_kata(p, subs)
        for sec in doc.sections:
            for area in (sec.header, sec.footer):
                for p in area.paragraphs:
                    _ganti_kata(p, subs)
        doc.save(str(docx))
    lengkapi_bab(docx, d, simpulan)
    pastikan_saipi(docx)
    final = docx.with_name(docx.name.replace("LHP-SUBSTANSI", awalan, 1))
    if final != docx:
        docx.replace(final)
    return final


def _ctx(d: Path) -> dict:
    out: dict = {}
    p = d / "context.md"
    if p.is_file():
        for b in p.read_text(encoding="utf-8").splitlines():
            m = re.match(r"^\s*[-*]?\s*(Kode|Obyek|Objek|Nomor ST|Tanggal ST)\s*:\s*(.+)$", b, re.I)
            if m:
                out[m.group(1).strip().lower().replace("objek", "obyek")] = m.group(2).strip()
            elif b.startswith("|"):
                sel = [c.strip() for c in b.strip().strip("|").split("|")]
                if len(sel) >= 2 and sel[0].lower() in ("kode", "objek", "obyek", "nomor st", "tanggal st"):
                    out.setdefault(sel[0].lower().replace("objek", "obyek"), sel[1])
    return out


def _butir(doc, teks: str) -> None:
    """Butir daftar; template ber-KOP tak selalu punya gaya 'List Bullet' — di server
    INTEGRAL memo konsultansi gagal dirender karena itu (KeyError)."""
    try:
        doc.add_paragraph(teks, style="List Bullet")
    except KeyError:
        doc.add_paragraph("• " + teks)


def _tujuan_lingkup(d: Path) -> tuple[str, str]:
    t = (d / "context.md").read_text(encoding="utf-8") if (d / "context.md").exists() else ""
    tj = re.search(r"^Tujuan\s*:\s*(.+?)(?=\n\n|\n##|\nRuang Lingkup)", t, re.M | re.S)
    rl = re.search(r"^Ruang Lingkup\s*:\s*(.+?)(?=\n\n|\n##|\Z)", t, re.M | re.S)
    return (tj.group(1).strip() if tj else ""), (rl.group(1).strip() if rl else "")


def _bab_wajib(doc, d: Path, a: dict, rekomendasi: list[str] | None = None) -> None:
    """Bab yang disyaratkan QC SAIPI (KOM-001..004) untuk laporan profil non-KKSA."""
    tj, rl = _tujuan_lingkup(d)
    doc.add_heading("Tujuan dan Sasaran", level=1); doc.add_paragraph(tj or "[DIISI AUDITOR — Tujuan dari DPP]")
    doc.add_heading("Ruang Lingkup", level=1); doc.add_paragraph(rl or "[DIISI AUDITOR — Ruang Lingkup dari DPP]")
    if rekomendasi is not None:
        doc.add_heading("Rekomendasi", level=1)
        for x in rekomendasi or ["[DIISI AUDITOR — rekomendasi]"]:
            _butir(doc, x)


def _kepala(doc, judul_besar: str, a: dict, ctx: dict, d: Path, dasar_label: str = "Dasar") -> None:
    doc.add_heading(judul_besar, level=0)
    doc.add_paragraph(a.get("judul") or judul_besar.title())
    meta = doc.add_paragraph()
    meta.add_run(f"Auditan: {a.get('auditi') or ctx.get('obyek', '-')}\n")
    meta.add_run(f"{dasar_label}: {a.get('dasar_permintaan') or ctx.get('nomor st', '-')}\n")
    meta.add_run(f"Kode penugasan: {ctx.get('kode', d.name)}")


def _isi_penanda(doc, peta: dict) -> None:
    """Ganti {{KUNCI}} per paragraf; nilai ber-baris-ganda disisipkan sebagai paragraf baru."""
    from copy import deepcopy
    from docx.text.paragraph import Paragraph
    for p in list(_paragraf_semua(doc)):
        if "{{" not in p.text:
            continue
        teks = p.text
        multi = None
        for k, v in peta.items():
            tag = "{{" + k + "}}"
            if tag not in teks:
                continue
            baris = [x for x in str(v).split("\n") if x.strip()] or ["[DIISI AUDITOR]"]
            if len(baris) > 1 and teks.strip() == tag:
                multi = baris
            else:
                teks = teks.replace(tag, " ".join(baris))
        if multi:
            _set(p, multi[0])
            acuan = p._p
            for x in multi[1:]:
                baru = deepcopy(p._p)
                acuan.addnext(baru); acuan = baru
                _set(Paragraph(baru, p._parent), x)
        else:
            _set(p, teks)


def render_memo(d: Path, a: dict, akar: Path) -> Path:
    """Memo konsultansi di atas kerangka resmi `template-lhp-konsultansi-umum.docx` (Nota Dinas +
    bab B–F). Server INTEGRAL memakai kerangka ini hanya sebagai KOP dan menulis memo di
    bawahnya — penanda {{B_PERTANYAAN}}…{{F_ASUMSI_BATASAN}} tertinggal kosong. Di sini diisi."""
    items = json.loads((d / "_LHP" / "saran.json").read_text(encoding="utf-8"))
    if not isinstance(items, list) or not items:
        raise ValueError("_LHP/saran.json kosong — isi [{pertanyaan, dasar_hukum[], telaah, pendapat, saran?, asumsi_batasan?}]")
    for i, it in enumerate(items, 1):
        kurang = [k for k in ("pertanyaan", "telaah", "pendapat") if not str(it.get(k) or "").strip()]
        if not it.get("dasar_hukum"):
            kurang.append("dasar_hukum")
        if kurang:
            raise ValueError(f"saran.json butir {i}: {', '.join(kurang)} kosong — pendapat tanpa telaah & dasar hukum tak boleh terbit")
    ctx = _ctx(d)
    kerangka = akar / "templates" / "_skeleton-lhp" / "template-lhp-konsultansi-umum.docx"
    doc = Document(str(kerangka))
    dh: list[str] = []
    for it in items:
        for x in it.get("dasar_hukum") or []:
            if x and x not in dh:
                dh.append(x)
    asumsi = [str(it["asumsi_batasan"]) for it in items if it.get("asumsi_batasan")] or [
        "Pendapat disusun berdasarkan dokumen dan informasi yang tersedia pada penugasan ini; bila fakta berbeda, pendapat dapat berubah."]
    t = (d / "context.md").read_text(encoding="utf-8") if (d / "context.md").exists() else ""
    ta = re.search(r"Tahun Anggaran\s*\|\s*([^|\n]+)", t)
    peta = {
        "B_PERTANYAAN": "\n".join(f"{i}. {it['pertanyaan']}" for i, it in enumerate(items, 1)),
        "C_DASAR_HUKUM": "\n".join(f"{i}. {x}" for i, x in enumerate(dh, 1)),
        "D_TELAAH": "\n".join(f"{i}. {it['telaah']}" for i, it in enumerate(items, 1)),
        "E_PENDAPAT": "\n".join(f"{i}. Pendapat: {it['pendapat']}" + (f" Saran: {it['saran']}" if it.get("saran") else "")
                                for i, it in enumerate(items, 1)),
        "F_ASUMSI_BATASAN": "\n".join(asumsi),
        "DASAR_PERMINTAAN": a.get("dasar_permintaan") or "[DIISI AUDITOR]",
        "NOMOR_ST": ctx.get("nomor st", "[DIISI AUDITOR]"), "TANGGAL_ST": ctx.get("tanggal st", "[DIISI AUDITOR]"),
        "NAMA_AUDITI": a.get("auditi") or ctx.get("obyek", "[DIISI AUDITOR]"),
        "TAHUN_ANGGARAN": ta.group(1).strip() if ta else "[DIISI AUDITOR]",
        "HAL_LHR": "Konsultansi " + (a.get("judul") or ctx.get("obyek", "")),
        "PENERIMA_LHP": "[DIISI AUDITOR]", "NOMOR_NOTA_DINAS": "[DIISI AUDITOR]", "TANGGAL_NOTA_DINAS": "[DIISI AUDITOR]",
        "LINK_SURVEI": "[DIISI AUDITOR]", "TEMBUSAN_LIST": "[DIISI AUDITOR]",
    }
    _isi_penanda(doc, peta)
    out = d / "_LHP" / "LHP-MEMO-KONSULTANSI.docx"  # QC mengenali awalan LHP/LHR/LHA/LHE/LP
    doc.save(str(out))
    lengkapi_bab(out, d, a.get("simpulan"))
    pastikan_saipi(out)
    return out


_RB_DIM = [("ketepatan", "Ketepatan Pelaksanaan"), ("ketercapaian", "Ketercapaian Output"),
           ("kualitas", "Kualitas Pelaksanaan"), ("kesesuaian", "Kesesuaian Waktu")]


def render_rb(d: Path, a: dict, akar: Path) -> Path:
    data = json.loads((d / "_LHP" / "penilaian-rb.json").read_text(encoding="utf-8"))
    komponen = data.get("komponen") if isinstance(data, dict) else None
    if not komponen:
        raise ValueError("_LHP/penilaian-rb.json: 'komponen' kosong")
    ctx = _ctx(d)
    doc = Document()
    _kepala(doc, "LAPORAN HASIL EVALUASI REFORMASI BIROKRASI", a, ctx, d)
    if a.get("gambaran_umum"):
        doc.add_heading("Gambaran Umum", level=1); doc.add_paragraph(a["gambaran_umum"])
    _bab_wajib(doc, d, a)
    doc.add_heading("Metodologi", level=1)
    doc.add_paragraph("Evaluasi dilaksanakan melalui penelaahan Rencana Aksi, realisasi, dan bukti dukung per komponen "
                      "pada empat dimensi penilaian. " + PERNYATAAN_SAIPI)
    doc.add_heading("Penilaian per Komponen Rencana Aksi (4 Dimensi)", level=1)
    t = doc.add_table(rows=1, cols=2 + len(_RB_DIM)); t.style = "Table Grid"
    h = t.rows[0].cells; h[0].text = "Komponen Renaksi"
    for j, (_, lbl) in enumerate(_RB_DIM, 1):
        h[j].text = lbl
    h[-1].text = "Catatan"
    for k in komponen:
        r = t.add_row().cells; r[0].text = str(k.get("nama", "-"))
        for j, (key, _) in enumerate(_RB_DIM, 1):
            r[j].text = str(k.get(key, "-"))
        r[-1].text = str(k.get("catatan", ""))
    if data.get("analisis_dampak"):
        doc.add_heading("Analisis Dampak", level=1); doc.add_paragraph(str(data["analisis_dampak"]))
    doc.add_heading("Simpulan", level=1)
    tak = sum(1 for k in komponen for key, _ in _RB_DIM if str(k.get(key, "")).lower().startswith("tidak"))
    doc.add_paragraph(a.get("simpulan") or f"Dari {len(komponen)} komponen Rencana Aksi yang dievaluasi, terdapat "
                      f"{tak} penilaian dimensi yang belum sesuai sebagaimana diuraikan pada tabel.")
    doc.add_heading("Rekomendasi", level=1)
    for x in [x for x in (data.get("aoi") or []) if x] or ["[DIISI AUDITOR — rekomendasi/AoI]"]:
        _butir(doc, str(x))
    out = d / "_LHP" / "LHE-RB.docx"
    doc.save(str(out))
    return out


def render_pendampingan(d: Path, a: dict, akar: Path) -> Path:
    items = json.loads((d / "_LHP" / "kegiatan-pendampingan.json").read_text(encoding="utf-8"))
    if not isinstance(items, list) or not items:
        raise ValueError("_LHP/kegiatan-pendampingan.json kosong — isi [{tanggal, jenis_kegiatan, deskripsi, hasil, …}]")
    ctx = _ctx(d)
    doc = Document()
    _kepala(doc, "LAPORAN HASIL PENDAMPINGAN PENGADAAN", a, ctx, d, "Dasar Penugasan")
    tgl = sorted(str(it.get("tanggal")) for it in items if it.get("tanggal"))
    if tgl:
        doc.add_paragraph(f"Periode Pendampingan: {tgl[0]}" + (f" s.d. {tgl[-1]}" if tgl[0] != tgl[-1] else ""))
    doc.add_paragraph("Catatan: laporan ini berisi rangkaian KEGIATAN PENDAMPINGAN yang telah diselesaikan tim "
                      "Inspektorat II atas permintaan unit kerja. Pendampingan bersifat advisory dan preventif — tidak "
                      "memberikan keyakinan dan tidak mengikat pejabat berwenang. " + PERNYATAAN_SAIPI)
    if a.get("gambaran_umum"):
        doc.add_heading("Gambaran Umum", level=1); doc.add_paragraph(a["gambaran_umum"])
    _bab_wajib(doc, d, a)
    doc.add_heading(f"I. Kegiatan Pendampingan yang Telah Diselesaikan ({len(items)})", level=1)
    t = doc.add_table(rows=1, cols=6); t.style = "Table Grid"
    for i, hd in enumerate(["No", "Tanggal", "Jenis Kegiatan", "Pihak Didampingi", "Deskripsi", "Hasil"]):
        t.rows[0].cells[i].text = hd
    for i, it in enumerate(items, 1):
        r = t.add_row().cells
        for j, v in enumerate([i, it.get("tanggal", "-"), it.get("jenis_kegiatan", "-"), it.get("pihak_didampingi", "-"),
                               it.get("deskripsi", "-"), it.get("hasil", "-")]):
            r[j].text = str(v)
    tl = [it for it in items if it.get("tindak_lanjut")]
    if tl:
        doc.add_heading(f"II. Hal yang Masih Memerlukan Tindak Lanjut ({len(tl)})", level=1)
        for i, it in enumerate(tl, 1):
            p = doc.add_paragraph(); p.add_run(f"{i}. {it.get('jenis_kegiatan', '-')} ({it.get('tanggal', '-')}): ").bold = True
            p.add_run(str(it["tindak_lanjut"]))
    doc.add_heading("Rekomendasi", level=1)
    for x in [str(it["tindak_lanjut"]) for it in tl] or ["Tidak terdapat hal yang memerlukan tindak lanjut."]:
        _butir(doc, x)
    doc.add_heading("III. Kesimpulan", level=1)
    doc.add_paragraph(a.get("kesimpulan") or (
        f"Tim Inspektorat II telah menyelesaikan {len(items)} kegiatan pendampingan pengadaan pada "
        f"{a.get('auditi') or ctx.get('obyek', 'unit kerja')}. Pendampingan ini tidak menggantikan kewenangan "
        f"PPK/PA/KPA atas keputusan pengadaan."))
    out = d / "_LHP" / "LHP-PENDAMPINGAN.docx"
    doc.save(str(out))
    return out


def sisa_isian(docx: Path) -> list[str]:
    """Bagian substansi yang masih '[DIISI — …]' (bukan penanda HITL [DIISI AUDITOR]/[DIISI AUDITI])."""
    doc = Document(str(docx))
    teks = "\n".join(p.text for p in _paragraf_semua(doc))
    return sorted({m for m in re.findall(r"\[DIISI[^\]]*\]", teks)
                   if not m.upper().startswith(("[DIISI AUDITOR", "[DIISI AUDITI")) and m.upper() != "[DIISI]"})
